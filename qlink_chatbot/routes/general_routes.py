import io
import os
import time
import uuid

from docx import Document
from fastapi import APIRouter, Body, File, Request, UploadFile, Header
from fastapi.responses import JSONResponse
from pymongo import MongoClient
from qlink_chatbot.utils.cloudflare_client import s3
from qlink_chatbot.utils.wa.send_sarthak_img import send_template_message

from qlink_chatbot.database.mongo_utils import (
    delete_alert_by_id,
    list_all_alerts,
    return_system_prompt,
    update_system_prompt,
    agent_login
)
from qlink_chatbot.database.pinecone_utils import (
    chunk_text,
    delete_record_by_id,
    fetch_records_with_metadata,
    get_record_by_id,
    list_records_by_label,
    store_vector_summary,
)
from qlink_chatbot.utils.logger_config import logger

MONGO_URI = os.getenv("MONGO_URI")
client = MongoClient(MONGO_URI)
db = client["JR"]
sessions_collection = db["users"]

general_router = APIRouter()

def short_id():
    return uuid.uuid4().hex[:8]

@general_router.get("/ping")
def ping():
    logger.info("Ping endpoint called")
    return {"message": "Jaipur Rugs web backend API is up and running"}


@general_router.get("/geo")
async def geo_check(request: Request, ip: str = ""):
    """Check what country/currency would be detected for an IP.

    Pass ?ip=1.2.3.4 to test a specific IP, or leave blank to use your own.
    """
    from qlink_chatbot.utils.geo_utils import get_geo
    target_ip = ip or request.client.host
    result = await get_geo(target_ip)
    return {"ip": target_ip, **result}

@general_router.post("/login")
async def agent_login_route(payload: dict):
    try:
        emp_id = payload.get("emp_id", "").strip().lower()
        password = payload.get("password", "").strip()

        if not emp_id or not password:
            return JSONResponse(
                {"message": "emp_id and password are required"},
                status_code=400
            )

        result = agent_login(emp_id=emp_id, password=password)

        if not result["success"]:
            return JSONResponse({"message": result["message"]}, status_code=401)

        return JSONResponse(
            {
                "message": "Login successful",
                "data": result["data"],
            },
            status_code=200
        )

    except Exception as e:
        logger.error("Login error", extra={"error": str(e)})
        return JSONResponse(
            {"message": "Server error"},
            status_code=500
        )


@general_router.post("/toggle/{session_id}")
def toggle_ai_mode(session_id: str):
    session = sessions_collection.find_one({"session_id": session_id})
    if not session:
        return JSONResponse({"error": "Session not found"}, status_code=404)

    new_status = not session.get("is_ai", True)
    sessions_collection.update_one(
        {"session_id": session_id},
        {"$set": {"is_ai": new_status}}
    )

    return {"session_id": session_id, "is_ai": new_status}


@general_router.get("/chat_history/{session_id}")
def get_chat_history(session_id: str):
    normalized_session_id = session_id.lower()
    session = sessions_collection.find_one({"session_id": normalized_session_id}, {"_id": 0})
    if not session:
        return {"session_id": normalized_session_id, "chat_history": []}
    return {"session_id": normalized_session_id, "chat_history": session.get("chat_history", [])}


@general_router.get("/users")
def get_all_users():
    try:
        users = []
        for user in sessions_collection.find(
            {},
            {
                "session_id": 1,
                "country_code": 1,
                "user_name": 1,
                "updated_at": 1
            }
        ).sort("updated_at", -1): 
            user["_id"] = str(user["_id"])
            users.append(user)
        logger.info(f"Fetched {len(users)} users successfully (sorted by latest updated).")
        return {"total_users": len(users), "users": users}
    except Exception as e:
        logger.error(f"Error fetching users: {e}")
        return JSONResponse({"error": "Failed to fetch users"}, status_code=500)


@general_router.get("/users/{session_id}")
def get_user_by_id(session_id: str):
    try:
        user = sessions_collection.find_one({"session_id": session_id})
        if not user:
            return JSONResponse({"error": "User not found"}, status_code=404)
        user["_id"] = str(user["_id"])
        return user
    except Exception as e:
        logger.error(f"Error fetching user: {e}")
        return JSONResponse({"error": "Failed to fetch user"}, status_code=500)
    
from pydantic import BaseModel

from qlink_chatbot.agent.stock_agent import openai_stock_response


class StockRequest(BaseModel):
    stock: str

@general_router.get("/kb/all/{lable}")
def list_all_agent_kb(lable: str):
    try:
        response = list_records_by_label(lable=lable)
        if not response:
            return JSONResponse({"error": "records not found"}, status_code=404)
        return response
    except Exception as e:
        logger.error(f"Error lisiting all records: {e}")
        return JSONResponse({"error": "Failed to list records"}, status_code=500)

@general_router.post("/kb/add")
async def add_agend_record_kb(data: dict = Body(...)):
    try:
        record = data.get("record")
        lable = data.get("lable", "agent")
        await store_vector_summary(
            session_id="self",
            summary=record,
            lable=lable
        )
        logger.info(f"Stored record in kb {record}")
        return True
    except Exception as e:
        logger.error(f"Error lisiting all records: {e}")
        return JSONResponse({"error": "Failed to store records"}, status_code=500)

@general_router.put("/kb/{id}")
async def update_record_from_kb(id: str, data: dict = Body(...)):
    try:
        record = data.get("record")
        lable = data.get("lable", "agent")
        delete_record_by_id(record_id=id)
        await store_vector_summary(
            session_id="self",
            summary=record,
            lable=lable
        )
        logger.info(f"Stored record in kb {record}")
        return True
    except Exception as e:
        logger.error(f"Error updating records: {e}")
        return JSONResponse({"error": "Failed to update records"}, status_code=500)

    
@general_router.get("/kb/{id}")
def get_record_from_kb(id: str):
    try:
        response = get_record_by_id(record_id=id)
        if not response:
            return JSONResponse({"error": "record not found"}, status_code=404)
        return response
    except Exception as e:
        logger.error(f"Error lisiting all records: {e}")
        return JSONResponse({"error": "Failed to get records"}, status_code=500)
    
@general_router.post("/kb/search")
async def search_records(data: dict = Body(...)):
    try:
        query = data.get("query")
        response = await fetch_records_with_metadata(query=query, top_k=15)
        if not response:
            return JSONResponse({"error": "record not found"}, status_code=404)
        return response
    except Exception as e:
        logger.error(f"Error lisiting all records: {e}")
        return JSONResponse({"error": "Failed to get records"}, status_code=500)

@general_router.delete("/kb/{id}")
def delete_record_from_kb(id: str):
    try:
        delete_record_by_id(record_id=id)
        return True
    except Exception as e:
        logger.error(f"Error lisiting all records: {e}")
        return JSONResponse({"error": "Failed to list records"}, status_code=500)

@general_router.post("/stock")
def get_stock_data(data: StockRequest):
    try:
        response = openai_stock_response(
            input=data.stock
        )

        return JSONResponse(
            content=response,
            status_code=200
        )
    
    except Exception as e:
        logger.error(f"Error fetching stocking request: {e}")
        return JSONResponse({"error": "Failed to fetch stock request"}, status_code=500)
    

@general_router.get("/system/prompt")
def get_system_prompt():
    """Fetch the current system prompt configuration."""
    try:
        response = return_system_prompt()
        if not response:
            return JSONResponse({"error": "System prompt not found"}, status_code=404)
        logger.info("Fetched system prompt successfully.")
        return response
    except Exception as e:
        logger.error("Error fetching system prompt", extra={"error": e})
        return JSONResponse({"error": "Failed to fetch system prompt"}, status_code=500)

@general_router.put("/system/prompt")
def update_system_prompt_route(data: dict = Body(...)):
    """Update the editable fields of the system prompt.
    Expected JSON body:
    {
      "system_identity": "...",
      "system_conversation_style": "...",
      "system_product_display_format": "...",
      "system_others": "..."
    }
    """
    try:
        system_identity = data.get("system_identity")
        system_conversation_style = data.get("system_conversation_style")
        system_product_display_format = data.get("system_product_display_format")
        system_others = data.get("system_others")

        result = update_system_prompt(
            system_identity=system_identity,
            system_conversation_style=system_conversation_style,
            system_product_display_format=system_product_display_format,
            system_others=system_others
        )

        return JSONResponse(content=result, status_code=200)
    except Exception as e:
        logger.error("Error updating system prompt", extra={"error": e})
        return JSONResponse({"error": "Failed to update system prompt"}, status_code=500)
    

@general_router.post("/kb/upload-docx")
async def upload_docx_and_store(file: UploadFile = File(...)):
    """Upload a .docx file, extract text, chunk it, and store chunks in vector DB.
    """
    try:
        if not file.filename.endswith(".docx"):
            return JSONResponse({"error": "Only .docx files are supported"}, status_code=400)

        content = await file.read()
        document = Document(io.BytesIO(content))

        full_text = "\n".join([para.text.strip() for para in document.paragraphs if para.text.strip()])
        if not full_text:
            return JSONResponse({"error": "No readable text found in document"}, status_code=400)

        chunks = chunk_text(full_text, max_length=1000, overlap=100)

        for i, chunk in enumerate(chunks):
            await store_vector_summary(
                session_id="self",
                summary=chunk,
                lable="general"
            )
            logger.info(f"Stored chunk {i+1}/{len(chunks)} in vector DB")

        return {"message": f"Uploaded {file.filename}, stored {len(chunks)} chunks in vector DB"}

    except Exception as e:
        logger.error(f"Error processing docx upload: {e}")
        return JSONResponse({"error": "Failed to process and store document"}, status_code=500)


@general_router.get("/alerts/all")
async def get_all_alerts():
    """Return all the agents alerts."""
    try:
        results = list_all_alerts()
        return results if results else []
    except Exception as e:
        logger.error("Error Fetching alerts", extra={"error": e})
        return JSONResponse({"error": "Failed to fetch all agent alerts"}, status_code=500)
    
@general_router.delete("/alerts/{id}")
async def delete_alert(id: str):
    """Route to delete the agent alert by id."""
    try:
        result = delete_alert_by_id(id=id)
        if not result:
            return JSONResponse({"message": "Alert not found"}, status_code=404)
        
        JSONResponse({"message": "Record deleted successfully"}, status_code=200)
    
    except Exception:
        JSONResponse({"message": "Error occured while deleting alert"}, status_code=500)
        logger.error("Error deleting record by id")

@general_router.get("/get-upload-url")
def get_upload_url(filename: str, email: str):
    email = email.lower()
    key = f"{email}/{short_id()}"

    try:
        url = s3.generate_presigned_url(
            ClientMethod="put_object",
            Params={
                "Bucket": "jr-chatbot", 
                "Key": key
            },
            ExpiresIn=60
        )

        public_url = f"https://pub-706af74a9f2443aa9e89918b8fd710b9.r2.dev/jr-chatbot/{key}"

        respone = {
            "upload_url": url,
            "final_url": public_url
        }

        # logger.error("Requst received to generate upload url", extra={"response": respone})
        return JSONResponse(respone, status_code=200)

    except Exception as e:
        logger.error("Error occured while generating upload url", extra={"error": e})

        return JSONResponse(
            {"message": "Error occured while generating upload url"}, 
            status_code=500
        )


class TriggerRequest(BaseModel):
    """Model for trigger."""
    phone_number: str

webhook_api = os.getenv("WEBHOOK_API")

@general_router.post("/wa/trigger")
async def dreamcast_webhook(
    request: TriggerRequest, x_api_key: str = Header(None)
):
    """Webhook message trigger."""
    try:
        if x_api_key != webhook_api:
            return JSONResponse(
                content={
                    "success": False,
                    "message": "Invalid or missing API Key",
                },
                status_code=401,
            )

        phone_number = request.phone_number
        logger.info("template trigger received", extra={"data": phone_number})
        send_template_message(
            phone_number=phone_number
        )

        return JSONResponse(
            content={
                "success": True
            },
            status_code=200,
        )

    except Exception as e:
        logger.exception(
            "Error in dreamcast webhook", extra={"exception": str(e)}
        )
        return JSONResponse(
            content={
                "success": False,
                "message": "Unexpected Error Occured",
                "error": str(e),
            },
            status_code=400,
        )


# ── Product Search REST endpoint ──────────────────────────────────────────────

@general_router.post("/search")
async def product_search(request: Request, payload: dict = Body(...)):
    """
    Unified product search endpoint for the frontend.

    Accepts structured filters directly — no keyword packing needed.

    Body (all fields optional):
    {
      "keyword":       "hand knotted red",   // free-text, parsed into filters
      "colors":        ["blue", "ivory"],
      "shapes":        ["round"],
      "sizes":         ["8x10"],
      "materials":     ["wool"],
      "constructions": ["hand knotted"],
      "styles":        ["traditional"],
      "price_max":     1000,
      "currency":      "USD",
      "weight_max":    8.0,
      "limit":         6
    }
    """
    from qlink_chatbot.utils.search_middleware import SearchFilters, search as _search

    try:
        keyword   = (payload.get("keyword") or "").strip()
        currency  = (payload.get("currency") or "INR").upper()
        limit     = min(int(payload.get("limit") or 6), 20)
        client_ip = request.client.host if request.client else ""

        if keyword:
            # Free-text path: parse keyword into filters then override with any
            # explicit params the caller also sent
            filters = SearchFilters.from_keyword(keyword, currency=currency, limit=limit)
            if payload.get("colors"):
                filters.colors = [c.lower() for c in payload["colors"]]
            if payload.get("shapes"):
                filters.shapes = [s.lower() for s in payload["shapes"]]
            if payload.get("sizes"):
                filters.sizes = payload["sizes"]
            if payload.get("materials"):
                filters.materials = [m.lower() for m in payload["materials"]]
            if payload.get("constructions"):
                filters.constructions = payload["constructions"]
            if payload.get("styles"):
                filters.styles = [s.lower() for s in payload["styles"]]
            if payload.get("price_max") is not None:
                filters.price_filter = {"currency": currency, "amount": float(payload["price_max"])}
            if payload.get("weight_max") is not None:
                filters.weight_filter = float(payload["weight_max"])
        else:
            # Structured path: explicit filter params only
            filters = SearchFilters.from_params(
                colors=payload.get("colors"),
                shapes=payload.get("shapes"),
                sizes=payload.get("sizes"),
                materials=payload.get("materials"),
                constructions=payload.get("constructions"),
                styles=payload.get("styles"),
                price_max=payload.get("price_max"),
                currency=currency,
                weight_max=payload.get("weight_max"),
                limit=limit,
            )

        results = await _search(filters, client_ip=client_ip)

        if isinstance(results, dict) and results.get("error"):
            return JSONResponse({"products": [], "total": 0, "message": results["error"]}, status_code=200)

        return JSONResponse({
            "products": results,
            "total": len(results),
            "source": "middleware",
            "filters_used": {
                "colors": filters.colors,
                "shapes": filters.shapes,
                "sizes": filters.sizes,
                "materials": filters.materials,
                "constructions": filters.constructions,
                "styles": filters.styles,
                "price_filter": filters.price_filter,
                "weight_filter": filters.weight_filter,
                "routed_to_mongodb": filters.needs_mongodb(),
            },
        }, status_code=200)

    except Exception as e:
        logger.error(f"Product search endpoint error: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)


# ── Visitor Insights endpoint ─────────────────────────────────────────────────

_INSIGHT_COLORS = {
    "red", "blue", "green", "yellow", "orange", "purple", "pink", "white",
    "black", "grey", "gray", "brown", "beige", "ivory", "cream", "navy",
    "teal", "turquoise", "gold", "silver", "rust", "coral", "indigo",
    "maroon", "burgundy", "charcoal", "olive", "mustard", "peach", "lavender",
}


@general_router.post("/visitor-insights/{session_id}")
def get_visitor_insights(session_id: str):
    """Return aggregated insights for a web visitor identified by session_id (email)."""
    try:
        sid = session_id.lower().strip()
        session = sessions_collection.find_one({"session_id": sid}, {"_id": 0})
        if not session:
            return JSONResponse({
                "session_id": sid,
                "found": False,
                "total_messages": 0,
                "total_searches": 0,
                "top_colors": [],
                "top_keywords": [],
                "chat_history": [],
                "previous_searches": [],
            }, status_code=200)

        chat_history = session.get("chat_history") or []
        previous_searches = session.get("previous_searches") or []

        # Count messages by role
        user_msgs = [m for m in chat_history if m.get("role") == "user"]

        # Extract colors from search filters
        color_counter: dict[str, int] = {}
        keyword_counter: dict[str, int] = {}
        for search in previous_searches:
            filters = search.get("filters") or {}
            for color in (filters.get("colors") or []):
                c = color.lower().strip()
                if c:
                    color_counter[c] = color_counter.get(c, 0) + 1
            kw = (search.get("keyword") or "").strip().lower()
            if kw:
                keyword_counter[kw] = keyword_counter.get(kw, 0) + 1
                for part in kw.split():
                    if part in _INSIGHT_COLORS:
                        color_counter[part] = color_counter.get(part, 0) + 1

        top_colors = sorted(color_counter.items(), key=lambda x: -x[1])
        top_keywords = sorted(keyword_counter.items(), key=lambda x: -x[1])

        return JSONResponse({
            "session_id": sid,
            "found": True,
            "user_name": session.get("user_name", ""),
            "country_code": session.get("country_code", ""),
            "last_active": str(session.get("updated_at", "")),
            "total_messages": len(chat_history),
            "user_messages": len(user_msgs),
            "total_searches": len(previous_searches),
            "top_colors": [{"color": c, "count": n} for c, n in top_colors[:5]],
            "top_keywords": [{"keyword": k, "count": n} for k, n in top_keywords[:5]],
            "chat_history": chat_history[-20:],
            "previous_searches": previous_searches[-10:],
        }, status_code=200)

    except Exception as e:
        logger.error(f"visitor-insights error for {session_id}: {e}")
        return JSONResponse({"error": str(e)}, status_code=500)
